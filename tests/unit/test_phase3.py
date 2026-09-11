"""第三阶段 FAQ、缺口、看板、生命周期和多格式解析安全回归。"""

import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from qdrant_client import QdrantClient
from sqlalchemy import select

from app.api.chat import get_chat_service
from app.core.business_dependencies import get_knowledge_service
from app.core.config import get_settings
from app.db.models import (
    FAQCandidate,
    FAQCluster,
    FAQEntry,
    IngestionJob,
    KnowledgeGap,
    KnowledgeTask,
    Message,
    MessageCitation,
    QuestionAudit,
)
from app.db.session import close_database, session_factory
from app.ingestion.readers import TextDocumentReader
from app.models.domain import EmbeddingBatch, SparseEmbedding
from app.repositories.authorized_qdrant import AuthorizedQdrantRepository
from app.services.authorized_knowledge import AuthorizedKnowledgeService
from app.services.chat import ChatService
from app.services.document_service import DocumentService
from app.services.faq import FAQService, cosine, normalize_question
from app.services.gaps import GapService
from app.services.identity_service import bootstrap_admin
from main import app


class Embedder:
    def embed(self, texts):
        vectors = [[1.0, 0.0, 0.0, 0.0] if "安装" in x or "Apollo" in x else [0.0, 1.0, 0.0, 0.0] for x in texts]
        return EmbeddingBatch(dense=vectors, sparse=[SparseEmbedding(indices=[1], values=[1.0]) for _ in texts])


class Reranker:
    def rerank(self, query, texts):
        return [0.9 for _ in texts]


class LLM:
    async def stream(self, messages):
        yield "依据资料可知 Apollo-77。[S1]"


class Phase3Test(unittest.TestCase):
    def setUp(self):
        close_database()
        self.tmp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {
            "DATABASE_URL": "sqlite:///" + self.tmp.name.replace("\\", "/") + "/test.db",
            "JWT_SECRET": "phase3-unit-only-secret-000000000000",
            "LOG_LEVEL": "ERROR", "FAQ_MIN_OCCURRENCES": "2",
        })
        self.env.start(); get_settings.cache_clear()
        self.client = TestClient(app, raise_server_exceptions=True); self.client.__enter__()
        self.vector = QdrantClient(":memory:")
        self.repo = AuthorizedQdrantRepository(self.vector, "phase3", 4)
        embed = Embedder()
        self.knowledge = AuthorizedKnowledgeService(
            DocumentService(reader=TextDocumentReader(1_000_000), embedder=embed, repository=self.repo,
                            max_chunk_size=1000, chunk_overlap=100, embedding_model_id="unit"),
            self.repo, embed, Reranker())
        app.dependency_overrides[get_knowledge_service] = lambda: self.knowledge
        app.dependency_overrides[get_chat_service] = lambda: ChatService(self.knowledge, LLM())
        with session_factory()() as db:
            admin = bootstrap_admin(db, "tenant-a", "admin", "Phase3-password-42")
            db.commit(); self.admin_id = admin.id
        self.headers = self.login("tenant-a", "admin")
        self.kb = self.request("POST", "/knowledge-bases", {"name": "研发库"}).json()["id"]

    def tearDown(self):
        app.dependency_overrides.clear(); self.client.__exit__(None, None, None)
        self.vector.close(); self.env.stop(); get_settings.cache_clear(); self.tmp.cleanup()

    def login(self, tenant, username):
        response = self.client.post("/auth/login", json={"tenant_id": tenant, "username": username, "password": "Phase3-password-42"})
        self.assertEqual(response.status_code, 200, response.text)
        return {"Authorization": "Bearer " + response.json()["access_token"]}

    def request(self, method, path, body=None, **kwargs):
        return self.client.request(method, path, json=body, headers=kwargs.pop("headers", self.headers), **kwargs)

    def upload(self, name="apollo.md", content=b"# Apollo\nApollo-77 is the internal code."):
        response = self.client.post("/documents/import", data={"kb_id": self.kb}, files={"file": (name, content)}, headers=self.headers)
        self.assertEqual(response.status_code, 201, response.text)
        doc_id = response.json()["id"]
        self.request("PUT", f"/documents/{doc_id}/permissions", {"entries": [{"permission_type": "global"}]})
        return doc_id

    def test_normalization_and_cosine(self):
        self.assertEqual(normalize_question("  Apollo 怎么安装？？ "), "apollo 怎么安装")
        self.assertAlmostEqual(cosine([1, 0], [1, 0]), 1.0)

    def test_all_document_parsers(self):
        reader = TextDocumentReader(1_000_000)
        self.assertIn("Markdown", reader.read_bytes(filename="a.md", data=b"# Markdown", tenant_id="t", kb_id="k").content)
        self.assertIn("TXT", reader.read_bytes(filename="a.txt", data=b"TXT", tenant_id="t", kb_id="k").content)
        docx_buffer = io.BytesIO(); docx = DocxDocument(); docx.add_heading("Word heading", 1); docx.add_paragraph("Word body"); docx.save(docx_buffer)
        self.assertIn("Word body", reader.read_bytes(filename="a.docx", data=docx_buffer.getvalue(), tenant_id="t", kb_id="k").content)
        # 直接用生产依赖 pypdf 构造带文本流的最小 PDF，避免测试环境额外依赖 ReportLab/Pillow。
        pdf_buffer = io.BytesIO(); writer = PdfWriter(); page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        })
        content = DecodedStreamObject(); content.set_data(b"BT /F1 12 Tf 72 720 Td (PDF body Apollo-77) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content); writer.write(pdf_buffer)
        self.assertIn("PDF body", reader.read_bytes(filename="a.pdf", data=pdf_buffer.getvalue(), tenant_id="t", kb_id="k").content)

    def test_batch_and_duplicate_import(self):
        response = self.client.post("/documents/import/batch", data={"kb_id": self.kb}, files=[
            ("files", ("a.md", b"# one")), ("files", ("b.exe", b"bad")),
        ], headers=self.headers)
        self.assertEqual((response.json()["success_count"], response.json()["failed_count"]), (1, 1))
        first = self.upload(); second = self.upload()
        self.assertEqual(first, second)

    def test_background_batch_validation_failure_does_not_leave_queued_jobs(self):
        with patch.object(get_settings(), "max_upload_bytes", 4):
            response = self.client.post(
                "/documents/import/jobs/batch",
                data={"kb_id": self.kb},
                files=[
                    ("files", ("valid.md", b"# A")),
                    ("files", ("oversized.txt", b"12345")),
                ],
                headers=self.headers,
            )

        self.assertEqual(response.status_code, 400, response.text)
        with session_factory()() as db:
            jobs = list(db.scalars(select(IngestionJob).where(IngestionJob.kb_id == self.kb)))
        self.assertEqual(jobs, [])

    def test_background_import_job_has_real_persisted_progress(self):
        created = self.client.post(
            "/documents/import/jobs",
            data={"kb_id": self.kb, "category": "异步导入"},
            files={"file": ("folder/async.md", b"# Async\nApollo background ingestion")},
            headers=self.headers,
        )
        self.assertEqual(created.status_code, 202, created.text)
        job_id = created.json()["id"]
        status = self.request("GET", f"/ingestion-jobs/{job_id}")
        self.assertEqual(status.status_code, 200, status.text)
        self.assertEqual((status.json()["status"], status.json()["progress"]), ("ready", 100))
        self.assertIsNotNone(status.json()["document_id"])
        with session_factory()() as db:
            self.assertEqual(db.get(IngestionJob, job_id).stage, "解析与向量索引完成")
        batch = self.client.post(
            "/documents/import/jobs/batch",
            data={"kb_id": self.kb},
            files=[("files", ("folder/a.md", b"# A")), ("files", ("folder/b.txt", b"B"))],
            headers=self.headers,
        )
        self.assertEqual(batch.status_code, 202, batch.text)
        statuses = [self.request("GET", f"/ingestion-jobs/{job['id']}").json() for job in batch.json()["items"]]
        self.assertEqual([job["status"] for job in statuses], ["ready", "ready"])

    def test_delete_removes_qdrant(self):
        doc_id = self.upload()
        self.assertGreater(self.repo.count_document(tenant_id="tenant-a", kb_id=self.kb, document_id=doc_id), 0)
        self.request("DELETE", f"/documents/{doc_id}")
        self.assertEqual(self.repo.count_document(tenant_id="tenant-a", kb_id=self.kb, document_id=doc_id), 0)

    def test_gap_aggregation_and_tenant_scope(self):
        gap = GapService().record("tenant-a", self.kb, "未知问题？", "empty_retrieval")
        GapService().record("tenant-a", self.kb, "未知问题?", "empty_retrieval")
        GapService().record("tenant-a", self.kb, "未知问题？", "low_rerank_score")
        result = self.request("GET", f"/knowledge-bases/{self.kb}/gaps")
        self.assertEqual(result.status_code, 200)
        rows = result.json()
        empty = next(row for row in rows if row["gap_type"] == "empty_retrieval")
        repeated = next(row for row in rows if row["gap_type"] == "repeated_unresolved")
        self.assertEqual(empty["occurrence_count"], 2)
        self.assertEqual(repeated["occurrence_count"], 3)
        self.assertEqual(self.request("POST", f"/gaps/{gap.id}/resolve").json()["status"], "resolved")

    def test_dashboard_uses_real_events(self):
        self.upload()
        self.request("POST", "/metrics/page-view", {"route": "/dashboard", "kb_id": self.kb})
        self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装？"})
        result = self.request("GET", f"/dashboard/summary?days=7&kb_id={self.kb}")
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["pv"], 1)
        self.assertEqual(result.json()["document_count"], 1)
        self.assertEqual(result.json()["question_count"], 1)
        self.assertEqual(result.json()["question_uv"], 1)
        self.assertEqual(result.json()["daily_question_count"], 1)
        self.assertEqual(result.json()["weekly_question_count"], 1)
        self.assertIn("knowledge_unit_count", result.json())
        self.assertIn("average_response_ms", result.json()["trend"][0])
        self.assertTrue(result.json()["top_questions"])

    def test_faq_real_mining_publish_and_acl_recheck(self):
        doc_id = self.upload()
        for _ in range(2):
            result = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装？"})
            self.assertEqual(result.status_code, 200, result.text)
        candidates = self.request("GET", f"/knowledge-bases/{self.kb}/faq/candidates").json()
        self.assertEqual(len(candidates), 1)
        self.assertGreater(candidates[0]["confidence_score"], 0)
        self.assertTrue(candidates[0]["question_samples"])
        self.assertEqual(candidates[0]["source_document_ids"], [doc_id])
        candidate_id = candidates[0]["id"]
        reviewed = self.request(
            "PATCH",
            f"/faq/candidates/{candidate_id}",
            {"status": "approved", "question": "Apollo 标准安装方法？", "answer": "Apollo-77。[S1]"},
        )
        self.assertEqual(reviewed.status_code, 200)
        self.assertEqual(reviewed.json()["standard_question"], "Apollo 标准安装方法？")
        published = self.request("POST", f"/faq/candidates/{candidate_id}/publish")
        self.assertEqual(published.status_code, 200)
        self.assertEqual(published.json()["question"], "Apollo 标准安装方法？")
        cached = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装"}).json()
        self.assertTrue(cached["faq_hit"])
        self.assertTrue(any(key[1:] == ("tenant-a", self.kb) for key in FAQService._cache))
        class NoEmbeddingAllowed:
            def embed(self, _):
                raise AssertionError("精确 FAQ 缓存命中不应调用远程 Embedding")

        original_embedder = self.knowledge.embedder
        self.knowledge.embedder = NoEmbeddingAllowed()
        exact = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 标准安装方法？"}).json()
        self.assertTrue(exact["faq_hit"])
        self.knowledge.embedder = original_embedder
        # 撤销文档 ACL 后，FAQ 不再命中，也不能返回缓存内容。
        self.request("PUT", f"/documents/{doc_id}/permissions", {"entries": []})
        denied = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装"}).json()
        self.assertFalse(denied["faq_hit"])
        self.assertNotIn("Apollo-77", denied["answer"])

    def test_citation_endpoint_rechecks_owner_and_acl(self):
        doc_id = self.upload()
        answer = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo"}).json()
        citation = answer["citations"][0]
        visible = self.request("GET", f"/citations/{answer['message_id']}/{citation['chunk_id']}")
        self.assertEqual(visible.status_code, 200)
        self.request("PUT", f"/documents/{doc_id}/permissions", {"entries": []})
        self.assertEqual(self.request("GET", f"/citations/{answer['message_id']}/{citation['chunk_id']}").status_code, 404)

    def test_sales_cannot_learn_from_faq_citation_or_dashboard(self):
        role = self.request("POST", "/roles", {"name": "sales-role", "permission_codes": [
            "knowledge:read", "chat:use", "dashboard:view"]}).json()["id"]
        sales_id = self.request("POST", "/users", {"username": "sales", "password": "Phase3-password-42", "role_ids": [role]}).json()["id"]
        self.request("POST", f"/knowledge-bases/{self.kb}/members", {"user_id": sales_id, "role": "viewer"})
        doc_id = self.upload()
        self.request("PUT", f"/documents/{doc_id}/permissions", {"entries": [{"permission_type": "user", "target_id": self.admin_id}]})
        for _ in range(2): self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装？"})
        candidate = self.request("GET", f"/knowledge-bases/{self.kb}/faq/candidates").json()[0]
        self.request("PATCH", f"/faq/candidates/{candidate['id']}", {"status": "approved"})
        self.request("POST", f"/faq/candidates/{candidate['id']}/publish")
        sales = self.login("tenant-a", "sales")
        self.assertEqual(self.request("GET", f"/knowledge-bases/{self.kb}/faqs", headers=sales).json(), [])
        answer = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装？"}, headers=sales).json()
        self.assertFalse(answer["faq_hit"]); self.assertNotIn("Apollo-77", answer["answer"])
        dashboard = self.request("GET", f"/dashboard/summary?kb_id={self.kb}", headers=sales).json()
        self.assertEqual(dashboard["document_count"], 0)

    def test_knowledge_unit_crud_chunk_settings_and_ledger(self):
        response = self.client.post(
            "/documents/import",
            data={"kb_id": self.kb, "category": "技术", "chunk_size": "256", "chunk_overlap": "32"},
            files={"file": ("settings.md", b"# Settings\nApollo setup")}, headers=self.headers,
        )
        self.assertEqual(response.status_code, 201, response.text)
        doc_id = response.json()["id"]
        ledger = self.request("GET", f"/documents?kb_id={self.kb}").json()[0]
        self.assertEqual((ledger["category"], ledger["chunk_size"], ledger["chunk_overlap"]), ("技术", 256, 32))
        self.assertTrue(ledger["permission_labels"])
        updated = self.request("PATCH", f"/documents/{doc_id}", {
            "title": "Apollo 安装说明", "category": "平台", "is_enabled": False,
        })
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertFalse(updated.json()["is_enabled"])
        self.request("PATCH", f"/documents/{doc_id}", {"is_enabled": True})
        reindexed = self.client.post(
            f"/documents/{doc_id}/reindex",
            data={"chunk_size": "128", "chunk_overlap": "16"},
            files={"file": ("settings.md", b"# Settings\nApollo adjusted chunk settings")},
            headers=self.headers,
        )
        self.assertEqual(reindexed.status_code, 200, reindexed.text)
        self.assertEqual((reindexed.json()["chunk_size"], reindexed.json()["chunk_overlap"]), (128, 16))
        self.assertEqual(self.request("DELETE", f"/documents/{doc_id}").status_code, 200)

    def test_restricted_prompt_and_question_audit(self):
        role = self.request("POST", "/roles", {"name": "visitor", "permission_codes": [
            "knowledge:read", "chat:use", "dashboard:view"]}).json()["id"]
        visitor_id = self.request("POST", "/users", {"username": "visitor", "password": "Phase3-password-42",
                                                      "role_ids": [role]}).json()["id"]
        self.request("POST", f"/knowledge-bases/{self.kb}/members", {"user_id": visitor_id, "role": "viewer"})
        doc_id = self.upload()
        self.request("PUT", f"/documents/{doc_id}/permissions", {
            "entries": [{"permission_type": "user", "target_id": self.admin_id}]
        })
        visitor = self.login("tenant-a", "visitor")
        result = self.request("POST", "/chat", {"kb_id": self.kb, "query": "Apollo 怎么安装？"}, headers=visitor)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertTrue(result.json()["restricted_sources_detected"])
        self.assertIn("无权查阅", result.json()["answer"])
        self.assertNotIn("Apollo-77", result.json()["answer"])
        audits = self.request("GET", f"/knowledge-bases/{self.kb}/question-audits").json()
        audit = next(x for x in audits if x["user_id"] == visitor_id)
        self.assertIn(doc_id, audit["recalled_document_ids"])
        self.assertIn(doc_id, audit["denied_document_ids"])
        self.assertEqual(audit["allowed_document_ids"], [])

    def test_gap_diagnostic_and_one_click_knowledge_task(self):
        gap = GapService().record(
            "tenant-a", self.kb, "海外清关延误怎么办？", "empty_retrieval", "补充清关说明",
            department_id="dept-logistics", highest_similarity_score=0.21, suggested_category="物流",
        )
        result = self.request("GET", f"/knowledge-bases/{self.kb}/gaps").json()[0]
        self.assertEqual(result["department_id"], "dept-logistics")
        self.assertEqual(result["suggested_category"], "物流")
        task = self.request("POST", f"/gaps/{gap.id}/knowledge-task")
        self.assertEqual(task.status_code, 201, task.text)
        with session_factory()() as db:
            self.assertIsNotNone(db.scalar(select(KnowledgeTask).where(KnowledgeTask.gap_id == gap.id)))

    def test_model_configuration_excludes_secrets_and_refreshes(self):
        updated = self.request("PATCH", "/system/model-config", {
            "bge_base_url": "http://127.0.0.1:18000",
            "reranker_base_url": "http://127.0.0.1:18000",
            "llm_base_url": "https://example.invalid/v1",
            "llm_model": "enterprise-model",
            "embedding_model_id": "bge-m3",
        })
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertEqual(updated.json()["llm_model"], "enterprise-model")
        self.assertIsInstance(updated.json()["llm_api_key_configured"], bool)
        self.assertNotIn("secret", updated.text.lower())

    def test_no_forbidden_placeholders(self):
        root = os.path.join(os.path.dirname(__file__), "..", "..", "app")
        offenders = []
        for base, _, files in os.walk(root):
            for name in files:
                if name.endswith(".py"):
                    text = Path(base, name).read_text(encoding="utf-8")
                    if "TODO" in text or "\n    pass\n" in text:
                        offenders.append(name)
        self.assertEqual(offenders, [])


if __name__ == "__main__":
    unittest.main()
