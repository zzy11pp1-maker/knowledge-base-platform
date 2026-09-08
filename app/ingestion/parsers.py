"""统一文档解析适配器；所有格式输出规范化文本后复用同一 RAG 流程。"""

from abc import ABC, abstractmethod
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.exceptions import DocumentReadError


class ParserAdapter(ABC):
    suffixes: frozenset[str]

    @abstractmethod
    def parse(self, data: bytes) -> tuple[str, str]:
        """返回规范化文本和解析器标识。"""


class TextParser(ParserAdapter):
    suffixes = frozenset({".md", ".markdown", ".txt"})

    def parse(self, data: bytes) -> tuple[str, str]:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return data.decode(encoding), f"text:{encoding}"
            except UnicodeDecodeError:
                continue
        raise DocumentReadError("文本编码无法识别，仅支持 UTF-8 或 GB18030")


class PDFParser(ParserAdapter):
    suffixes = frozenset({".pdf"})

    def parse(self, data: bytes) -> tuple[str, str]:
        if not data.startswith(b"%PDF"):
            raise DocumentReadError("PDF 文件签名不正确")
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise DocumentReadError("暂不支持加密 PDF")
            pages = []
            for index, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append(f"# 第 {index} 页\n\n{text}")
        except DocumentReadError:
            raise
        except Exception as exc:
            raise DocumentReadError("PDF 解析失败") from exc
        if not pages:
            raise DocumentReadError("PDF 没有可提取文字；扫描件需要后续 OCR 能力")
        return "\n\n".join(pages), "pypdf"


class DOCXParser(ParserAdapter):
    suffixes = frozenset({".docx"})

    def parse(self, data: bytes) -> tuple[str, str]:
        if not data.startswith(b"PK"):
            raise DocumentReadError("DOCX 文件签名不正确")
        try:
            document = DocxDocument(BytesIO(data))
            blocks = []
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style = (paragraph.style.name or "").lower()
                if style.startswith("heading"):
                    number = "".join(ch for ch in style if ch.isdigit()) or "2"
                    blocks.append("#" * min(int(number), 6) + " " + text)
                else:
                    blocks.append(text)
            for table in document.tables:
                for row in table.rows:
                    values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(values):
                        blocks.append(" | ".join(values))
        except Exception as exc:
            raise DocumentReadError("DOCX 解析失败") from exc
        if not blocks:
            raise DocumentReadError("DOCX 没有可提取文字")
        return "\n\n".join(blocks), "python-docx"


class ParserRegistry:
    def __init__(self, adapters=None):
        adapters = adapters or [TextParser(), PDFParser(), DOCXParser()]
        self._by_suffix = {suffix: adapter for adapter in adapters for suffix in adapter.suffixes}

    @property
    def supported_suffixes(self):
        return frozenset(self._by_suffix)

    def parse(self, filename: str, data: bytes) -> tuple[str, str]:
        suffix = Path(filename).suffix.lower()
        adapter = self._by_suffix.get(suffix)
        if adapter is None:
            raise DocumentReadError("仅支持 PDF、DOCX、Markdown 和 TXT 文件")
        return adapter.parse(data)
